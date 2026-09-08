"""Artificial vehicles through production routes and real SQLite/PostgreSQL storage.

No production data, model response or evidence pack is fabricated here.
"""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models.vehicle import KnowledgeItem
from app.schemas.vehicle import DRIVE_TYPES, ENERGY_TYPES, SERIES_OPTIONS
from app.services.vehicle_service import EXCEL_HEADERS
from test_vehicles_api import ADMIN_HEADERS, client, setup_function, _create_vehicle

FIELDS = dict(energy_type="range_extended", series="国产新能源", displacement="1.5L",
              battery_capacity_kwh="82.500000000000000000000000001", drive_type="four_wheel_drive")


def details(code):
    with SessionLocal() as db:
        return deepcopy(db.scalar(select(KnowledgeItem).where(KnowledgeItem.item_id == code)).payload)


def workbook_bytes(rows, *, version=None):
    template = client.get("/api/vehicles/excel/template", headers=ADMIN_HEADERS)
    assert template.status_code == 200
    book = load_workbook(BytesIO(template.content))
    if version:
        book["_元数据"]["B1"] = version
    for row in rows:
        book["车辆信息"].append([row.get(header) for header in EXCEL_HEADERS])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def preview(content):
    return client.post("/api/vehicles/excel/preview", files={"file": ("vehicles.xlsx", content)}, headers=ADMIN_HEADERS)


def confirm(result):
    return client.post(f"/api/vehicles/excel/{result['preview_id']}/confirm", headers=ADMIN_HEADERS)


@pytest.mark.parametrize("field,values", [("energy_type", ENERGY_TYPES), ("series", SERIES_OPTIONS), ("drive_type", DRIVE_TYPES)])
def test_all_selections_roundtrip_search_preserve_clear(field, values):
    for value in values:
        car = _create_vehicle(**{**FIELDS, field: value})
        code = car["vehicle_code"]
        assert car[field] == value
        assert details(code)["data"]["additional_details"][field] == value
        updated = client.put(f"/api/vehicles/{code}", json={"public_price": "12.88"}, headers=ADMIN_HEADERS)
        assert updated.status_code == 200
        assert updated.json()["data"][field] == value
        keyword = ENERGY_TYPES.get(value, DRIVE_TYPES.get(value, value))
        found = client.get("/api/vehicles", params={"keyword": keyword}, headers=ADMIN_HEADERS).json()["data"]
        assert code in [item["vehicle_code"] for item in found["items"]]
        assert client.put(f"/api/vehicles/{code}", json={field: None}, headers=ADMIN_HEADERS).json()["data"][field] is None
        assert field not in details(code)["data"]["additional_details"]


@pytest.mark.parametrize("field,value", [
    ("energy_type", "新能源纯电"), ("energy_type", ""), ("energy_type", "unknown"),
    ("series", "卡罗拉"), ("series", ""), ("drive_type", "四驱"), ("drive_type", "all_wheel_drive"),
    ("displacement", "x" * 101), ("displacement", 1.5),
    *[("battery_capacity_kwh", value) for value in ["-1", "0", "0.00", "1e2", "1E2", "75 kWh", "NaN", "Infinity", "abc", "1,000", "9" * 101, 82.5, True, {}, ["75"]]],
])
def test_invalid_values_rejected_without_touching_saved_vehicle(field, value):
    code = _create_vehicle(**FIELDS)["vehicle_code"]
    before = details(code)
    for method, url in (("POST", "/api/vehicles"), ("PUT", f"/api/vehicles/{code}")):
        response = client.request(method, url, json={"display_name": "错误不应写入", field: value}, headers=ADMIN_HEADERS)
        assert response.status_code == 400, response.text
        assert response.json()["code"] == "VALIDATION_ERROR"
    assert details(code) == before


@pytest.mark.parametrize("capacity", ["75", "82.5", "00075.500", "0." + "0" * 97 + "1", "9" * 100])
def test_capacity_exact_decimal_text_without_rounding(capacity):
    car = _create_vehicle(**{**FIELDS, "battery_capacity_kwh": f"  {capacity}  ", "displacement": "  任意排量原文  "})
    assert car["battery_capacity_kwh"] == capacity
    assert car["displacement"] == "任意排量原文"
    assert details(car["vehicle_code"])["data"]["additional_details"]["battery_capacity_kwh"] == capacity
    cleared = client.put(f"/api/vehicles/{car['vehicle_code']}", json={"battery_capacity_kwh": "  ", "displacement": ""}, headers=ADMIN_HEADERS).json()["data"]
    assert cleared["battery_capacity_kwh"] is None and cleared["displacement"] is None


def test_legacy_series_survives_unrelated_update_and_explicit_selection_replaces_it():
    code = _create_vehicle(**FIELDS, vin="TEST-INTERNAL", internal_notes="保留内部资料")["vehicle_code"]
    with SessionLocal() as db:
        row = db.scalar(select(KnowledgeItem).where(KnowledgeItem.item_id == code))
        payload = deepcopy(row.payload)
        payload["data"]["additional_details"]["series"] = "卡罗拉"
        row.payload = payload  # Explicit old database fixture, not an allowed new API input.
        db.commit()
    response = client.put(f"/api/vehicles/{code}", json={"public_price": "13.00"}, headers=ADMIN_HEADERS)
    assert response.status_code == 200, response.text
    updated = response.json()["data"]
    assert updated["series"] == "卡罗拉"
    assert updated["vin"] == "TEST-INTERNAL"
    for field, value in FIELDS.items():
        if field != "series": assert updated[field] == value
    assert "原车系：卡罗拉" in details(code)["data"]["specs"]
    assert client.get("/api/vehicles", params={"keyword": "卡罗拉"}, headers=ADMIN_HEADERS).json()["data"]["total"] == 1
    for series in ("日系车", None):
        assert client.put(f"/api/vehicles/{code}", json={"series": series}, headers=ADMIN_HEADERS).json()["data"]["series"] == series


def test_missing_fields_remain_unknown_and_do_not_change_listing_rules():
    car = client.post("/api/vehicles", json={"display_name": "只有名称"}, headers=ADMIN_HEADERS).json()["data"]
    assert all(car[field] is None for field in FIELDS)
    assert details(car["vehicle_code"])["data"]["specs"] == ""
    assert client.post(f"/api/vehicles/{car['vehicle_code']}/list", headers=ADMIN_HEADERS).status_code == 409


def test_excel_all_choices_precise_capacity_preview_no_write_blank_update_preserves():
    rows = [{"展示名称": f"导入-{index}", "车系": series, "能源类型": list(ENERGY_TYPES.values())[index % 4],
             "排量": " 1.5L ", "电池包容量（kWh）": FIELDS["battery_capacity_kwh"],
             "驱动方式": list(DRIVE_TYPES.values())[index % 3]} for index, series in enumerate(SERIES_OPTIONS)]
    content = workbook_bytes(rows)
    book = load_workbook(BytesIO(content))
    assert book["_元数据"]["B1"].value == "1.2"
    assert [cell.value for cell in book["车辆信息"][1]] == EXCEL_HEADERS
    assert len(book["车辆信息"].data_validations.dataValidation) == 3
    result = preview(content).json()["data"]
    assert result["can_confirm"], result
    with SessionLocal() as db: assert db.scalar(select(KnowledgeItem)) is None
    assert confirm(result).status_code == 200
    for index, row in enumerate(result["rows"]):
        code = row["vehicle_code"]
        old = client.get(f"/api/vehicles/{code}", headers=ADMIN_HEADERS).json()["data"]
        assert old["energy_type"] == list(ENERGY_TYPES)[index % 4]
        assert old["drive_type"] == list(DRIVE_TYPES)[index % 3]
        assert old["series"] == SERIES_OPTIONS[index]
        assert old["displacement"] == "1.5L"
        assert old["battery_capacity_kwh"] == FIELDS["battery_capacity_kwh"]
        change = preview(workbook_bytes([{"车辆编号": code, "公开售价": "22.00"}])).json()["data"]
        assert confirm(change).status_code == 200
        new = client.get(f"/api/vehicles/{code}", headers=ADMIN_HEADERS).json()["data"]
        assert {key: new[key] for key in FIELDS} == {key: old[key] for key in FIELDS}


@pytest.mark.parametrize("column,value", [("能源类型", "electric"), ("车系", "卡罗拉"), ("驱动方式", "全时四驱"), ("排量", "x" * 101), *[("电池包容量（kWh）", value) for value in ["1e2", "82.5kWh", 0, -2, "9" * 101, "=75+5"]]])
def test_excel_invalid_row_blocks_entire_batch(column, value):
    result = preview(workbook_bytes([{"展示名称": "合法"}, {"展示名称": "错误", column: value}])).json()["data"]
    assert result["error_count"] == 1, result
    assert result["rows"][1]["row_number"] == 3 and result["rows"][1]["errors"]
    assert confirm(result).status_code == 409
    with SessionLocal() as db: assert db.scalar(select(KnowledgeItem)) is None


@pytest.mark.parametrize("raw", ["82.5", "82.500000000000000000000000001", "1e2"])
def test_excel_numeric_xml_never_roundtrips_through_float(raw):
    data = workbook_bytes([{"展示名称": "数值单元格", "电池包容量（kWh）": 82.5}])
    rewritten = BytesIO()
    with ZipFile(BytesIO(data)) as src, ZipFile(rewritten, "w") as dst:
        for name in src.namelist():
            part = src.read(name)
            if name == "xl/worksheets/sheet1.xml":
                assert b'<c r="H2" t="n"><v>82.5</v></c>' in part
                part = part.replace(b'<c r="H2" t="n"><v>82.5</v></c>', f'<c r="H2" t="n"><v>{raw}</v></c>'.encode())
            dst.writestr(name, part)
    result = preview(rewritten.getvalue()).json()["data"]
    if "e" in raw:
        assert not result["can_confirm"]
    else:
        assert result["can_confirm"]
        assert confirm(result).status_code == 200
        assert details(result["rows"][0]["vehicle_code"])["data"]["additional_details"]["battery_capacity_kwh"] == raw


def test_old_excel_template_reports_download_latest():
    result = preview(workbook_bytes([], version="1.1"))
    assert result.status_code == 409
    assert result.json()["data"]["expected"] == "1.2"


def test_excel_trailing_formula_cannot_disappear_as_an_empty_row():
    result = preview(workbook_bytes([{"展示名称": "合法"}, {"电池包容量（kWh）": "=75+5"}])).json()["data"]
    assert result["total_rows"] == 2
    assert result["error_count"] == 1
    assert confirm(result).status_code == 409
